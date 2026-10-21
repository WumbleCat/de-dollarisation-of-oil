"""Look up rows of the project's data tracker: data/Oil_Dedollarisation_Data_Tracker.xlsx, sheet "Data Sources".

    python .claude/skills/data-sources/scripts/tracker_sources.py                 # one line per row (all 50)
    python .claude/skills/data-sources/scripts/tracker_sources.py D13 D14         # full record for these IDs
    python .claude/skills/data-sources/scripts/tracker_sources.py -c "Market technicals"
    python .claude/skills/data-sources/scripts/tracker_sources.py -c geopolitics --json
    python .claude/skills/data-sources/scripts/tracker_sources.py --folders       # category -> ./data folder map
    python .claude/skills/data-sources/scripts/tracker_sources.py --grep fred     # rows whose text mentions "fred"

Every row maps to a folder: ./data/<category slug>/ where the slug is the Category column lower-cased with
spaces replaced by underscores (Fundamentals -> data/fundamentals, Market technicals -> data/market_technicals ...).
Needs openpyxl (in the project venv: .venv/Scripts/python.exe). The workbook is only ever READ here.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import textwrap
from pathlib import Path

TRACKER_REL = Path("data") / "Oil_Dedollarisation_Data_Tracker.xlsx"
SHEET = "Data Sources"

# Columns to show in the one-line listing, in order.
LIST_COLS = ["ID", "Category", "Variable", "Specific series / ticker", "Primary source", "Access", "Format / API", "Status"]
# Column order for full records (anything else in the sheet is appended after these).
RECORD_COLS = ["ID", "Category", "Variable", "Specific series / ticker", "Primary source", "URL",
               "Alternative / backup source", "Frequency", "Coverage from", "Access", "Format / API", "Notes",
               "Status", "Date retrieved", "Saved file / location", "Findings / limitations"]


def category_slug(category: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(category).strip().lower()).strip("_")


def find_tracker(explicit: str | None) -> Path:
    """The workbook path: --xlsx, $TRACKER_XLSX, or the nearest ancestor of cwd / this script holding data/<xlsx>."""
    if explicit:
        p = Path(explicit).expanduser()
        if p.exists():
            return p
        sys.exit(f"tracker workbook not found: {p}")
    env = os.environ.get("TRACKER_XLSX")
    if env and Path(env).exists():
        return Path(env)
    starts = [Path.cwd(), Path(__file__).resolve().parent]
    for start in starts:
        for d in [start, *start.parents]:
            cand = d / TRACKER_REL
            if cand.exists():
                return cand
    sys.exit(f"could not find {TRACKER_REL.as_posix()} above {Path.cwd()} - pass --xlsx <path> or set TRACKER_XLSX")


def load_rows(xlsx: Path, sheet: str) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: run with the project venv python (.venv/Scripts/python.exe)")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"sheet {sheet!r} not in {xlsx.name}; sheets: {wb.sheetnames}")
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for raw in it:
        rec = {h: (v if v is not None else "") for h, v in zip(header, raw) if h}
        if not str(rec.get("ID", "")).strip():
            continue
        for k, v in rec.items():
            if hasattr(v, "isoformat"):
                rec[k] = v.isoformat()[:10]
            elif not isinstance(v, str):
                rec[k] = str(v)
            else:
                rec[k] = v.strip()
        rec["folder"] = f"data/{category_slug(rec.get('Category', ''))}"
        rows.append(rec)
    wb.close()
    return rows


def matches(rec: dict, ids: list[str], category: str | None, grep: str | None) -> bool:
    if ids and rec["ID"].upper() not in ids:
        return False
    if category:
        want = category_slug(category)
        if want not in (category_slug(rec["Category"]), rec["folder"].split("/")[-1]):
            return False
    if grep:
        blob = " ".join(str(v) for v in rec.values()).lower()
        if grep.lower() not in blob:
            return False
    return True


def print_list(rows: list[dict]) -> None:
    for r in rows:
        print(" | ".join(str(r.get(c, "")) for c in LIST_COLS) + f" | {r['folder']}")


def print_records(rows: list[dict]) -> None:
    wrap = textwrap.TextWrapper(width=110, subsequent_indent=" " * 30)
    for r in rows:
        print(f"=== {r['ID']}  ({r['Category']} -> {r['folder']}) ===")
        cols = [c for c in RECORD_COLS if c in r] + [c for c in r if c not in RECORD_COLS and c != "folder"]
        for c in cols:
            v = r.get(c, "")
            if v == "":
                continue
            print(wrap.fill(f"{c:<28}: {v}"))
        print()


def print_folders(rows: list[dict]) -> None:
    seen: dict[str, list[str]] = {}
    for r in rows:
        seen.setdefault(f"{r['Category']} -> {r['folder']}", []).append(r["ID"])
    for k, ids in seen.items():
        print(f"{k:<45} {ids[0]}..{ids[-1]} ({len(ids)} rows)")


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ids", nargs="*", help="tracker IDs such as D13 (full record for each)")
    ap.add_argument("-c", "--category", help="filter by Category (name or folder slug, case-insensitive)")
    ap.add_argument("-g", "--grep", help="keep rows whose text contains this (case-insensitive)")
    ap.add_argument("--full", action="store_true", help="full records even when no IDs are given")
    ap.add_argument("--json", action="store_true", help="emit JSON instead of text")
    ap.add_argument("--folders", action="store_true", help="print the category -> ./data folder map and exit")
    ap.add_argument("--xlsx", help="path to the tracker workbook (default: data/Oil_Dedollarisation_Data_Tracker.xlsx)")
    ap.add_argument("--sheet", default=SHEET)
    a = ap.parse_args(argv)

    rows = load_rows(find_tracker(a.xlsx), a.sheet)
    if a.folders:
        print_folders(rows)
        return
    ids = [i.upper() for i in a.ids]
    sel = [r for r in rows if matches(r, ids, a.category, a.grep)]
    missing = [i for i in ids if i not in {r["ID"].upper() for r in sel}]
    if missing:
        print(f"warning: no row for {', '.join(missing)}", file=sys.stderr)
    if a.json:
        print(json.dumps(sel, indent=2, ensure_ascii=False))
    elif ids or a.full:
        print_records(sel)
    else:
        print_list(sel)
    if not sel:
        sys.exit(1)


if __name__ == "__main__":
    main()
